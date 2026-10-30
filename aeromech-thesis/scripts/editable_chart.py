# -*- coding: utf-8 -*-
"""Generate charts from supplied values, including a native PowerPoint chart."""
import math
import os
import json

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import figure_style as ST
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Cm


def render(spec, output_base):
    kind = spec.get('type')
    categories = spec.get('categories') or []
    series = spec.get('series') or []
    if kind not in ('bar', 'line') or not categories or not series:
        raise ValueError('Chart type, categories, and series are required')
    for entry in series:
        values = entry.get('values') or []
        if not entry.get('name') or len(values) != len(categories) or any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
            raise ValueError('Each series requires one finite number per category')
    data = CategoryChartData()
    data.categories = [str(x) for x in categories]
    for entry in series:
        data.add_series(entry['name'], entry['values'])
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(18), Cm(11)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED if kind == 'bar' else XL_CHART_TYPE.LINE_MARKERS
    chart = slide.shapes.add_chart(chart_type, Cm(0.7), Cm(0.5), Cm(16.6), Cm(9.8), data).chart
    if spec.get('title'):
        chart.has_title = True
        chart.chart_title.text_frame.text = str(spec['title'])
    chart.has_legend = len(series) > 1
    if spec.get('ylabel'):
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = str(spec['ylabel'])
    plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
    fig, ax = plt.subplots(figsize=(7.1, 4.3), constrained_layout=True)
    fig.get_layout_engine().set(w_pad=0.22, h_pad=0.22)
    colors = [ST.COLOR_PALETTE[key] for key in ('primary', 'secondary', 'accent', 'risk', 'success')]
    ax.set_axisbelow(True)
    ax.grid(axis='y', color='#DEE5E4', linewidth=0.8)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.spines['left'].set_color('#AAB8B6')
    ax.spines['bottom'].set_color('#AAB8B6')
    ax.tick_params(colors='#334946', length=0, pad=6)
    positions = list(range(len(categories)))
    for index, entry in enumerate(series):
        if kind == 'bar':
            width = 0.55 / len(series)
            ax.bar([x - 0.275 + width * (index + 0.5) for x in positions], entry['values'], width=width * 0.85, label=entry['name'], color=colors[index % len(colors)])
        else:
            ax.plot(positions, entry['values'], marker='o', linewidth=2, label=entry['name'], color=colors[index % len(colors)])
    ax.set_xticks(positions, [str(x) for x in categories])
    ax.set_title(str(spec.get('title') or ''))
    if spec.get('ylabel'):
        ax.set_ylabel(spec['ylabel'])
    legend = ax.legend(fontsize=12) if len(series) > 1 else None
    for label in list(ax.get_xticklabels()) + list(ax.get_yticklabels()) + [ax.title, ax.yaxis.label]:
        label.set_fontsize(12)
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    scale = 2.54 / fig.dpi
    def bbox(artist):
        bb = artist.get_window_extent(renderer=renderer)
        return [bb.x0 * scale, bb.y0 * scale, bb.x1 * scale, bb.y1 * scale]
    low, high = ax.get_ylim()
    visible_y = [label for value, label in zip(ax.get_yticks(), ax.get_yticklabels())
                 if low <= value <= high and label.get_visible() and label.get_text()]
    ylabels = [{'bbox': bbox(t)} for t in visible_y]
    xlabels = [bbox(t) for t in ax.get_xticklabels() if t.get_visible() and t.get_text()]
    bars = [{'bbox': bbox(p), 'color': colors[i % len(colors)]}
            for i, container in enumerate(ax.containers) for p in container.patches]
    lines = [{'id': str(i), 'color': colors[i % len(colors)]} for i, _ in enumerate(ax.lines)]
    texts = []
    artists = list(ax.get_xticklabels()) + visible_y + [ax.title, ax.yaxis.label]
    if legend:
        artists += list(legend.get_texts())
    for artist in artists:
        if artist.get_visible() and artist.get_text():
            bounds = bbox(artist)
            texts.append({'box': '__note__', 'x0': bounds[0], 'y0': bounds[1],
                          'x1': bounds[2], 'y1': bounds[3], 'fs': artist.get_fontsize()})
    os.makedirs(os.path.dirname(os.path.abspath(output_base)), exist_ok=True)
    result = {kind: output_base + '.' + kind for kind in ('pptx', 'png', 'svg')}
    result['layout'] = output_base + '.layout.json'
    with open(result['layout'], 'w', encoding='utf-8') as f:
        json.dump({'name': os.path.basename(output_base), 'figure_kind': 'editable_chart',
                   'pixel_width': int(7.1 * 300), 'pixel_height': int(4.3 * 300), 'w_cm': 7.1 * 2.54,
                   'h_cm': 4.3 * 2.54, 'dpi': 300, 'boxes': [], 'edges': [],
                   'texts': texts, 'min_font_pt': min(t['fs'] for t in texts),
                   'axes_rect': bbox(ax), 'bars': bars, 'lines': lines,
                   'y_labels': ylabels, 'x_labels': xlabels, 'value_labels': []},
                  f, ensure_ascii=False, indent=2)
    deck.save(result['pptx'])
    fig.savefig(result['png'], dpi=300, facecolor='white')
    fig.savefig(result['svg'], facecolor='white')
    plt.close(fig)
    return result
