# -*- coding: utf-8 -*-
"""Render a supplied FMEA table or risk matrix as editable cells and images."""
import os
import textwrap
import json

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import figure_style as ST
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Cm, Pt


def render(spec, output_base):
    columns = spec.get('columns') or []
    rows = spec.get('rows') or []
    roles = spec.get('roles')
    if not columns or not rows or any(not isinstance(row, list) or len(row) != len(columns) for row in rows):
        raise ValueError('Each row must contain one value per column')
    if roles is not None and (len(roles) != len(rows) or any(len(row) != len(columns) or any(role not in ('normal', 'medium', 'high') for role in row) for row in roles)):
        raise ValueError('Risk roles must match the data grid')
    data = [[str(c) for c in columns]] + [[str(c) for c in row] for row in rows]
    if any(len(value) > 40 for row in data for value in row):
        raise ValueError('Cell text too long; split the table')
    data = [[textwrap.fill(value, width=10) for value in row] for row in data]
    width = max(14, len(columns) * 3)
    row_lines = [max(value.count('\n') + 1 for value in row) for row in data]
    row_height = max(0.9, max(row_lines) * 0.5 + 0.3)
    height = max(4, len(data) * row_height + 1)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(width), Cm(height)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(len(data), len(columns), Cm(0.4), Cm(0.4), Cm(width - 0.8), Cm(height - 0.8)).table
    fills = {'normal': 'E8EEF3', 'medium': 'E7D2BD', 'high': 'E5CFC9'}
    colors = []
    for i, row in enumerate(data):
        color_row = []
        for j, value in enumerate(row):
            fill = '33607F' if i == 0 else fills[roles[i - 1][j] if roles else 'normal']
            cell = table.cell(i, j)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fill)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.font.name = 'Microsoft YaHei'
                paragraph.font.size = Pt(11)
                paragraph.font.color.rgb = RGBColor.from_string('FFFFFF' if i == 0 else '1F2933')
            color_row.append('#' + fill)
        colors.append(color_row)
    plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
    fig, ax = plt.subplots(figsize=(width / 2.54, height / 2.54))
    ax.axis('off')
    grid = ax.table(cellText=data, cellColours=colors, cellLoc='center', loc='center', bbox=[0, 0, 1, 1])
    grid.auto_set_font_size(False)
    grid.set_fontsize(11)
    for (i, _), cell in grid.get_celld().items():
        cell.set_edgecolor('#FFFFFF')
        cell.get_text().set_color('#FFFFFF' if i == 0 else ST.COLOR_PALETTE['text'])
    fig.subplots_adjust(left=0.05, right=0.95, top=0.86, bottom=0.14)
    os.makedirs(os.path.dirname(os.path.abspath(output_base)), exist_ok=True)
    result = {kind: output_base + '.' + kind for kind in ('pptx', 'png', 'svg')}
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    scale = 2.54 / fig.dpi
    boxes, texts = [], []
    for (i, j), cell in grid.get_celld().items():
        bb = cell.get_window_extent(renderer=renderer)
        tb = cell.get_text().get_window_extent(renderer=renderer)
        ident = f'{i}-{j}'
        boxes.append({'id': ident, 'x': bb.x0 * scale, 'y': bb.y0 * scale,
                      'w': bb.width * scale, 'h': bb.height * scale,
                      'fill': colors[i][j], 'stroke': '#FFFFFF',
                      'text_color': '#FFFFFF' if i == 0 else ST.COLOR_PALETTE['text']})
        texts.append({'box': ident, 'x0': tb.x0 * scale, 'y0': tb.y0 * scale,
                      'x1': tb.x1 * scale, 'y1': tb.y1 * scale, 'fs': 11})
    result['layout'] = output_base + '.layout.json'
    with open(result['layout'], 'w', encoding='utf-8') as f:
        json.dump({'name': os.path.basename(output_base), 'figure_kind': 'table',
                   'pixel_width': int(width / 2.54 * 300), 'pixel_height': int(height / 2.54 * 300),
                   'w_cm': width, 'h_cm': height, 'dpi': 300, 'boxes': boxes,
                   'texts': texts, 'edges': [], 'min_font_pt': 11},
                  f, ensure_ascii=False, indent=2)
    deck.save(result['pptx'])
    fig.savefig(result['png'], dpi=300, facecolor='white')
    fig.savefig(result['svg'], facecolor='white')
    plt.close(fig)
    return result
