# -*- coding: utf-8 -*-
"""Render explicitly positioned diagram nodes to editable slides and images."""
import os
import textwrap
import math
import json

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import figure_style as ST
from matplotlib.patches import Ellipse, FancyArrowPatch, FancyBboxPatch, Polygon
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Cm, Pt


def render(spec, output_base):
    nodes = spec.get('nodes') or []
    edges = spec.get('edges') or []
    ids = [str(n.get('id') or '') for n in nodes]
    if not ids or any(not i for i in ids) or len(set(ids)) != len(ids):
        raise ValueError('Nodes require unique IDs')
    by_id = dict(zip(ids, nodes))
    for node in nodes:
        if not node.get('text') or any(isinstance(node.get(k), bool) or not isinstance(node.get(k), (int, float)) or not math.isfinite(node[k]) for k in ('x', 'y', 'w', 'h')) or node['w'] <= 0 or node['h'] <= 0 or node['x'] < 0 or node['y'] < 0:
            raise ValueError('Node needs text and positive dimensions')
        if node.get('shape', 'rounded') not in ('rounded', 'rectangle', 'diamond', 'ellipse') or node.get('role', 'normal') not in ('normal', 'medium', 'high'):
            raise ValueError('Unsupported node shape or role')
        if node.get('semantic') and node['semantic'] not in ('top_event', 'gate', 'basic_event'):
            raise ValueError('Unsupported semantic role')
        width_chars = max(1, int(node['w'] * 2.2))
        lines = textwrap.wrap(str(node['text']), width=width_chars, replace_whitespace=False) or ['']
        max_font = min(12, node['h'] * 28.35 / (len(lines) * 1.35))
        choices = [size for size in (12, 11.5, 11) if size <= max_font]
        if not choices:
            raise ValueError('Node text does not fit; enlarge the node')
        font_size = choices[0]
        node['_render_lines'] = '\n'.join(lines)
        node['_font_size'] = font_size
    if any(e.get('from') not in by_id or e.get('to') not in by_id for e in edges):
        raise ValueError('Edge references missing node')
    width = max(n['x'] + n['w'] for n in nodes) + 1
    height = max(n['y'] + n['h'] for n in nodes) + 1
    deck = Presentation()
    deck.slide_width, deck.slide_height = Cm(width), Cm(height)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    fig, ax = plt.subplots(figsize=(width / 2.54, height / 2.54))
    fig.subplots_adjust(left=0, right=1, bottom=0, top=1)
    ax.set_xlim(0, width)
    ax.set_ylim(height, 0)
    ax.axis('off')
    plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
    measured_texts = []
    boxes = []
    layout_edges = []
    for node in nodes:
        x, y, w, h = (node[k] for k in ('x', 'y', 'w', 'h'))
        kind = node.get('shape', 'rounded')
        role = node.get('role', 'normal')
        fill, stroke = {'normal': ('E8EEF3', '33607F'), 'medium': ('E7D2BD', 'B5663D'), 'high': ('E5CFC9', '9E4337')}[role]
        if node.get('semantic'):
            border, background, _ = ST.SEMANTIC_COLORS[node['semantic']]
            fill = ST.COLOR_PALETTE[background][1:]
            stroke = ST.COLOR_PALETTE[border][1:]
        ppt_shapes = {'rounded': MSO_SHAPE.ROUNDED_RECTANGLE, 'rectangle': MSO_SHAPE.RECTANGLE,
                      'diamond': MSO_SHAPE.DIAMOND, 'ellipse': MSO_SHAPE.OVAL}
        shape = slide.shapes.add_shape(ppt_shapes[kind], Cm(x), Cm(y), Cm(w), Cm(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape._element.spPr.append(OxmlElement('a:effectLst'))
        shape.text_frame.text = node['_render_lines']
        shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            paragraph.font.size = Pt(node['_font_size'])
            paragraph.font.name = 'Microsoft YaHei'
            paragraph.font.color.rgb = RGBColor.from_string('1F2933')
        if kind == 'ellipse':
            patch = Ellipse((x + w / 2, y + h / 2), w, h, facecolor='#' + fill, edgecolor='#' + stroke)
        elif kind == 'diamond':
            patch = Polygon([(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)], facecolor='#' + fill, edgecolor='#' + stroke)
        else:
            box = 'round,pad=0.02,rounding_size=0.1' if kind == 'rounded' else 'square,pad=0'
            patch = FancyBboxPatch((x, y), w, h, boxstyle=box, facecolor='#' + fill, edgecolor='#' + stroke)
        ax.add_patch(patch)
        text_artist = ax.text(x + w / 2, y + h / 2, node['_render_lines'], ha='center', va='center', fontsize=node['_font_size'], color=ST.COLOR_PALETTE['text'])
        measured_texts.append((node['id'], text_artist, node['_font_size']))
        boxes.append({'id': node['id'], 'x': x, 'y': height - y - h, 'w': w, 'h': h,
                      'fill': '#' + fill, 'stroke': '#' + stroke, 'text_color': ST.COLOR_PALETTE['text'],
                      'role': node.get('semantic'), 'text': str(node['text'])})
    for edge in edges:
        a, b = by_id[edge['from']], by_id[edge['to']]
        ac = (a['x'] + a['w'] / 2, a['y'] + a['h'] / 2)
        bc = (b['x'] + b['w'] / 2, b['y'] + b['h'] / 2)
        dx, dy = bc[0] - ac[0], bc[1] - ac[1]
        if dx == 0 and dy == 0:
            raise ValueError('Connected nodes must have distinct centers')
        def boundary(node, center, vx, vy):
            scales = []
            if vx:
                scales.append(node['w'] / (2 * abs(vx)))
            if vy:
                scales.append(node['h'] / (2 * abs(vy)))
            scale = min(scales)
            return center[0] + vx * scale, center[1] + vy * scale
        start = boundary(a, ac, dx, dy)
        end = boundary(b, bc, -dx, -dy)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(start[0]), Cm(start[1]), Cm(end[0]), Cm(end[1]))
        line.line.color.rgb = RGBColor.from_string('315F61')
        arrow = OxmlElement('a:tailEnd')
        arrow.set('type', 'triangle')
        line.line._get_or_add_ln().append(arrow)
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle='->', mutation_scale=12, color='#315F61'))
        layout_edges.append({'id': str(len(layout_edges) + 1),
                             'pts': [[start[0], height - start[1]], [end[0], height - end[1]]]})
        if edge.get('label'):
            mx, my = (start[0] + end[0]) / 2, (start[1] + end[1]) / 2
            label = slide.shapes.add_textbox(Cm(mx + 0.1), Cm(my - 0.25), Cm(2), Cm(0.5))
            label.text_frame.text = str(edge['label'])
            label.text_frame.paragraphs[0].font.size = Pt(11.5)
            label.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string('1F2933')
            edge_text = ax.text(mx + 0.1, my, str(edge['label']), va='center', fontsize=11.5, color='#1F2933')
            measured_texts.append(('__note__', edge_text, 11.5))
    os.makedirs(os.path.dirname(os.path.abspath(output_base)), exist_ok=True)
    result = {kind: output_base + '.' + kind for kind in ('pptx', 'png', 'svg')}
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    inverse = ax.transData.inverted()
    texts = []
    for node_id, artist, font_size in measured_texts:
        extent = artist.get_window_extent(renderer=renderer)
        x0, y0 = inverse.transform((extent.x0, extent.y0))
        x1, y1 = inverse.transform((extent.x1, extent.y1))
        texts.append({'box': node_id, 'x0': x0, 'x1': x1, 'y0': height - y0,
                      'y1': height - y1, 'fs': font_size})
    result['layout'] = output_base + '.layout.json'
    with open(result['layout'], 'w', encoding='utf-8') as f:
        json.dump({'name': os.path.basename(output_base), 'w_cm': width, 'h_cm': height,
                   'dpi': 300, 'figure_kind': 'editable_diagram',
                   'pixel_width': int(width / 2.54 * 300), 'pixel_height': int(height / 2.54 * 300),
                   'boxes': boxes, 'edges': layout_edges, 'texts': texts,
                   'min_font_pt': min(t['fs'] for t in texts)}, f, ensure_ascii=False, indent=2)
    deck.save(result['pptx'])
    fig.savefig(result['png'], dpi=300, facecolor='white')
    fig.savefig(result['svg'], facecolor='white')
    plt.close(fig)
    return result
