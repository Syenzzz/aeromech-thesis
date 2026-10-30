import os
import sys
import tempfile
import unittest

from pptx import Presentation
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'scripts'))
import editable_figure
import editable_chart
import editable_table


class FigureTypeTests(unittest.TestCase):
    def test_research_diagrams_preserve_supplied_labels_and_edges(self):
        examples = [
            ('技术路线', ['研究问题', '方法选择', '分析结论']),
            ('故障诊断', ['故障现象', '检查压力', '排除故障']),
            ('故障树', ['顶事件', 'OR门', '基础事件']),
            ('事件树', ['初始事件', '保护成功', '后果']),
        ]
        with tempfile.TemporaryDirectory() as root:
            for name, labels in examples:
                nodes = [{'id': str(i), 'text': label, 'x': 1, 'y': i * 2 + 1, 'w': 4, 'h': 1}
                         for i, label in enumerate(labels)]
                spec = {'nodes': nodes, 'edges': [{'from': '0', 'to': '1'}, {'from': '1', 'to': '2'}]}
                files = editable_figure.render(spec, os.path.join(root, name))
                shapes = Presentation(files['pptx']).slides[0].shapes
                self.assertEqual([shape.text for shape in shapes if shape.has_text_frame], labels)
                self.assertEqual(sum(1 for shape in shapes if shape.shape_type == 9), 2)

    def test_statistical_data_and_fmea_rows_survive(self):
        with tempfile.TemporaryDirectory() as root:
            for kind in ('bar', 'line'):
                files = editable_chart.render({'type': kind, 'categories': ['A', 'B'],
                                               'series': [{'name': '数据', 'values': [2, 4]}]},
                                              os.path.join(root, kind))
                chart = next(s.chart for s in Presentation(files['pptx']).slides[0].shapes if s.has_chart)
                self.assertEqual(list(chart.series[0].values), [2.0, 4.0])
            files = editable_table.render({'columns': ['故障模式', '严重度'], 'rows': [['泄漏', '高']]},
                                          os.path.join(root, 'fmea'))
            table = next(s.table for s in Presentation(files['pptx']).slides[0].shapes if s.has_table)
            self.assertEqual(table.cell(1, 0).text, '泄漏')


if __name__ == '__main__':
    unittest.main()
