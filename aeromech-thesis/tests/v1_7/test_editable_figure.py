import os
import sys
import tempfile
import unittest
import json

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'scripts'))
import editable_figure
from pptx import Presentation


class EditableFigureTests(unittest.TestCase):
    def test_diagram_outputs_editable_shapes_and_vector(self):
        spec = {'nodes': [{'id': 'a', 'text': 'Start', 'x': 1, 'y': 1, 'w': 3, 'h': 1},
                          {'id': 'b', 'text': 'Finish', 'x': 1, 'y': 3, 'w': 3, 'h': 1}],
                'edges': [{'from': 'a', 'to': 'b'}]}
        with tempfile.TemporaryDirectory() as root:
            files = editable_figure.render(spec, os.path.join(root, 'figure'))
            self.assertTrue(all(os.path.isfile(p) for p in files.values()))
            shapes = Presentation(files['pptx']).slides[0].shapes
            self.assertGreaterEqual(len(shapes), 3)
            self.assertEqual([s.text for s in shapes if s.has_text_frame], ['Start', 'Finish'])
            self.assertEqual(str(shapes[0].text_frame.paragraphs[0].font.color.rgb), '1F2933')
            with open(files['layout'], encoding='utf-8') as f:
                layout = json.load(f)
            self.assertEqual(len(layout['boxes']), 2)
            self.assertEqual(len(layout['edges']), 1)
            self.assertEqual(len(layout['texts']), 2)

    def test_missing_edge_target_fails(self):
        with tempfile.TemporaryDirectory() as root:
            with self.assertRaises(ValueError):
                editable_figure.render({'nodes': [], 'edges': [{'from': 'a', 'to': 'b'}]}, os.path.join(root, 'f'))

    def test_decision_and_edge_label_are_editable(self):
        spec = {'nodes': [{'id': 'a', 'text': 'Decision', 'x': 1, 'y': 1, 'w': 3, 'h': 2, 'shape': 'diamond'},
                          {'id': 'b', 'text': 'High risk', 'x': 1, 'y': 4, 'w': 4, 'h': 2, 'role': 'high'}],
                'edges': [{'from': 'a', 'to': 'b', 'label': 'Yes'}]}
        with tempfile.TemporaryDirectory() as root:
            files = editable_figure.render(spec, os.path.join(root, 'figure'))
            shapes = Presentation(files['pptx']).slides[0].shapes
            self.assertIn('Yes', [s.text for s in shapes if s.has_text_frame])


if __name__ == '__main__':
    unittest.main()
