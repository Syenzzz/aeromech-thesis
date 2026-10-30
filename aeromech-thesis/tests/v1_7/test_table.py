import os
import sys
import tempfile
import unittest
import json

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'scripts'))
import editable_table
import figure_visual_qa
from pptx import Presentation


class TableTests(unittest.TestCase):
    def test_table_data_remains_editable(self):
        spec = {'columns': ['Failure', 'Severity'], 'rows': [['Leak', 'High'], ['Wear', 'Medium']]}
        with tempfile.TemporaryDirectory() as root:
            files = editable_table.render(spec, os.path.join(root, 'figure'))
            slide = Presentation(files['pptx']).slides[0]
            table = next(shape.table for shape in slide.shapes if shape.has_table)
            self.assertEqual(table.cell(1, 0).text, 'Leak')
            self.assertEqual(table.cell(2, 1).text, 'Medium')
            self.assertTrue(os.path.exists(files['svg']))
            with open(files['layout'], encoding='utf-8') as f:
                metadata = json.load(f)
            self.assertEqual(metadata['figure_kind'], 'table')
            self.assertEqual(len(metadata['boxes']), 6)
            report = figure_visual_qa.run(root, out=root)
            self.assertFalse([i for i in report.items if i['status'] == 'FAIL'])

    def test_inconsistent_rows_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            with self.assertRaises(ValueError):
                editable_table.render({'columns': ['A', 'B'], 'rows': [['X']]}, os.path.join(root, 'x'))

    def test_risk_matrix_preserves_severity_colors(self):
        spec = {'columns': ['Probability', 'Low', 'High'],
                'rows': [['Frequent', 'Medium', 'High']],
                'roles': [['normal', 'medium', 'high']]}
        with tempfile.TemporaryDirectory() as root:
            files = editable_table.render(spec, os.path.join(root, 'risk'))
            table = next(s.table for s in Presentation(files['pptx']).slides[0].shapes if s.has_table)
            self.assertEqual(str(table.cell(1, 2).fill.fore_color.rgb), 'E5CFC9')


if __name__ == '__main__':
    unittest.main()
