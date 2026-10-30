import os
import sys
import tempfile
import unittest
import json

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', 'scripts'))
import editable_chart
import figure_visual_qa
from pptx import Presentation


class ChartTests(unittest.TestCase):
    def test_chart_contains_editable_data(self):
        spec = {'type': 'bar', 'title': 'Faults', 'categories': ['A', 'B'],
                'series': [{'name': 'Count', 'values': [3, 5]}]}
        with tempfile.TemporaryDirectory() as root:
            files = editable_chart.render(spec, os.path.join(root, 'figure'))
            self.assertTrue(all(os.path.isfile(p) for p in files.values()))
            slide = Presentation(files['pptx']).slides[0]
            chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
            self.assertEqual(list(chart.series[0].values), [3.0, 5.0])
            with open(files['layout'], encoding='utf-8') as f:
                metadata = json.load(f)
            self.assertEqual(len(metadata['bars']), 2)
            self.assertTrue(metadata['y_labels'])
            report = figure_visual_qa.run(root, out=root)
            self.assertFalse([i for i in report.items if i['status'] == 'FAIL'])

    def test_missing_data_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            with self.assertRaises(ValueError):
                editable_chart.render({'type': 'bar'}, os.path.join(root, 'figure'))


if __name__ == '__main__':
    unittest.main()
